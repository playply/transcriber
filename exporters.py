from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document


def _stamp(seconds: float | int | None) -> str:
    value = max(0, int(seconds or 0))
    hours, remainder = divmod(value, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _normalized_segment(segment: dict[str, Any]) -> dict[str, Any]:
    return {
        "start": segment.get("start"),
        "end": segment.get("end"),
        "speaker": segment.get("speaker", "UNKNOWN"),
        "resolved_name": segment.get("resolved_name"),
        "text": (segment.get("text") or "").strip(),
        "words": segment.get("words", []),
    }


def _paths(source: Path) -> dict[str, Path]:
    base = source.with_suffix("")
    return {
        "json": Path(f"{base}_transcript.json"),
        "txt": Path(f"{base}_transcript.txt"),
        "docx": Path(f"{base}_transcript.docx"),
    }


def _display_speaker(segment: dict[str, Any]) -> str:
    return str(segment.get("resolved_name") or segment.get("speaker") or "UNKNOWN")


def write_canonical_transcript(canonical: dict[str, Any], source: Path) -> dict[str, Path]:
    """Persist canonical JSON and regenerate TXT/DOCX without retranscription."""
    source = source.expanduser().resolve()
    paths = _paths(source)
    segments = [_normalized_segment(segment) for segment in canonical.get("segments", [])]
    canonical = dict(canonical)
    canonical["segments"] = segments

    tmp_json = paths["json"].with_suffix(paths["json"].suffix + ".tmp")
    tmp_json.write_text(
        json.dumps(canonical, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    tmp_json.replace(paths["json"])

    lines: list[str] = []
    for segment in segments:
        lines.append(
            f"[{_stamp(segment['start'])}–{_stamp(segment['end'])}] "
            f"{_display_speaker(segment)}: {segment['text']}"
        )
    paths["txt"].write_text("\n".join(lines) + "\n", encoding="utf-8")

    document = Document()
    document.add_heading(source.stem, level=1)
    for segment in segments:
        paragraph = document.add_paragraph()
        paragraph.add_run(
            f"[{_stamp(segment['start'])}–{_stamp(segment['end'])}] "
            f"{_display_speaker(segment)}: "
        ).bold = True
        paragraph.add_run(segment["text"])
    document.save(paths["docx"])

    return paths


def export_transcript(result: dict[str, Any], source: Path) -> dict[str, Path]:
    segments = [_normalized_segment(segment) for segment in result.get("segments", [])]
    canonical = {
        "source": {
            "filename": source.name,
            "path": str(source),
        },
        "processed_at": result.get("processed_at"),
        "language": result.get("language"),
        "models": result.get("models", {}),
        "speaker_count": result.get("speaker_count", 0),
        "detected_speakers": result.get("detected_speakers", []),
        "segments": segments,
    }
    return write_canonical_transcript(canonical, source)
